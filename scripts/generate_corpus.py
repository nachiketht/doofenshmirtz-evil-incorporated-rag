"""Reproducibly generate the Doofenshmirtz Evil Inc document corpus.

    pip install -e ".[corpus]"
    python scripts/generate_corpus.py            # writes docs/ + docs/manifest.json
    python scripts/generate_corpus.py --check    # verify docs/ is up to date

The seven original hand-written documents are never touched; they only get
manifest entries. Every generated file follows the filename contract
``Doofenshmirtz Evil Inc - <Name> v<X.Y>.<pdf|docx|md>`` and the numbered-heading
contract, and is read back with ``rag.reader`` to prove it parses into exactly
the headings it was generated from. Output is byte-for-byte deterministic
(fixed PDF/DOCX timestamps), so incremental ingest only re-embeds real changes.
"""

import argparse
import io
import json
import re
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "src"))

from scripts.corpus import FAMILIES
from scripts.corpus.existing import ORIGINALS

COMPANY = "Doofenshmirtz Evil Incorporated"
PREFIX = "Doofenshmirtz Evil Inc - "
BANNER = "TOP SECRET — EYES OF HEINZ DOOFENSHMIRTZ ONLY"
FIXED_ZIP_TIME = (2024, 1, 1, 0, 0, 0)
DECIMAL_AFTER_WORD = re.compile(r"(\S) (\d+\.\d+)")


# ------------------------------------------------------------- versions
def resolve(family: dict) -> list[dict]:
    """Expand derived versions into full section lists."""
    resolved: dict[str, dict] = {}
    for spec in family["versions"]:
        if "base" in spec:
            base = resolved[spec["base"]]
            sections = [list(item) for item in base["sections"]]
            for title in spec.get("remove", ()):
                before = len(sections)
                sections = [item for item in sections if item[0] != title]
                if len(sections) == before:
                    raise ValueError(f"{family['name']}: cannot remove {title!r}")
            for title, new in spec.get("replace", {}).items():
                index = next(
                    (i for i, item in enumerate(sections) if item[0] == title), None
                )
                if index is None:
                    raise ValueError(f"{family['name']}: cannot replace {title!r}")
                sections[index] = list(new)
            if "purpose" in spec:
                sections[0] = ["Purpose", spec["purpose"]]
            if "acknowledgment" in spec:
                sections[-1] = ["Acknowledgment", spec["acknowledgment"]]
            tail = sections.pop()  # Acknowledgment stays last
            sections.extend(list(item) for item in spec.get("add", ()))
            sections.append(tail)
        else:
            sections = [list(item) for item in spec["sections"]]
        framed = sections[0][0] == "Purpose" and sections[-1][0] == "Acknowledgment"
        if family.get("doc_type", "policy") == "policy" and not framed:
            raise ValueError(f"{family['name']} v{spec['version']}: bad frame")
        resolved[spec["version"]] = {**spec, "sections": sections}
    return [resolved[spec["version"]] for spec in family["versions"]]


def filename(family: dict, version: dict) -> str:
    return f"{PREFIX}{family['name']} v{version['version']}.{version['format']}"


def display_title(family: dict, version: dict) -> str:
    return f"{family.get('title', family['name'])} — Version {version['version']}"


def expected_headings(sections) -> list[str]:
    heads = []
    for number, (title, body) in enumerate(sections, start=1):
        heads.append(f"{number}. {title}")
        if isinstance(body, list):
            for sub, (subtitle, _text) in enumerate(body, start=1):
                heads.append(f"{number}.{sub} {subtitle}")
    return heads


# --------------------------------------------------------------- items
def items(sections):
    """Flatten sections into renderable items.

    Yields ("heading", text) | ("para", text) | ("table", rows) | ("bullet", text).
    """
    for number, (title, body) in enumerate(sections, start=1):
        yield "heading", f"{number}. {title}"
        if isinstance(body, str):
            yield "para", body
        elif isinstance(body, list):
            for sub, (subtitle, text) in enumerate(body, start=1):
                yield "para", f"{number}.{sub} {subtitle}. {text}"
        else:
            if body.get("intro"):
                yield "para", body["intro"]
            if body.get("table"):
                yield "table", body["table"]
            for bullet in body.get("bullets", ()):
                yield "bullet", bullet
            if body.get("after"):
                yield "para", body["after"]


def header_lines(family: dict, version: dict) -> list[str]:
    lines = [COMPANY, display_title(family, version)]
    if family.get("classification") == "top-secret":
        lines.append(BANNER)
    return lines


# ------------------------------------------------------------ renderers
def render_md(family, version) -> bytes:
    head = header_lines(family, version)
    out = [f"# {head[0]}", f"## {head[1]}"]
    out.extend(f"**{line}**" for line in head[2:])
    out.append("")
    for kind, value in items(version["sections"]):
        if kind == "heading":
            out.extend([f"## {value}", ""])
        elif kind == "para":
            match = re.match(r"^(\d+\.\d+ [^.]+\.)(.*)$", value)
            text = f"**{match.group(1)}**{match.group(2)}" if match else value
            out.extend([text, ""])
        elif kind == "bullet":
            out.append(f"- {value}")
        else:
            header, *rows = value
            out.append("| " + " | ".join(header) + " |")
            out.append("| " + " | ".join("---" for _ in header) + " |")
            out.extend("| " + " | ".join(row) + " |" for row in rows)
            out.append("")
    text = "\n".join(out).rstrip() + "\n"
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.encode("utf-8")


def _normalize_zip(raw: bytes) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(raw))
    target = io.BytesIO()
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as out:
        for info in sorted(source.infolist(), key=lambda item: item.filename):
            fixed = zipfile.ZipInfo(info.filename, date_time=FIXED_ZIP_TIME)
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = 0o600 << 16
            out.writestr(fixed, source.read(info.filename))
    return target.getvalue()


def render_docx(family, version) -> bytes:
    from datetime import UTC, datetime

    import docx

    document = docx.Document()
    props = document.core_properties
    props.author = "Doofenshmirtz Evil Inc Compliance (Norm)"
    props.title = display_title(family, version)
    props.created = props.modified = datetime(2024, 1, 1, tzinfo=UTC)
    props.last_modified_by = "Norm"
    props.revision = 1
    for line in header_lines(family, version):
        document.add_paragraph(line)
    for kind, value in items(version["sections"]):
        if kind == "heading":
            document.add_heading(value, level=2)
        elif kind == "para":
            document.add_paragraph(value)
        elif kind == "bullet":
            document.add_paragraph(f"- {value}")
        else:
            header, *rows = value
            table = document.add_table(rows=len(rows) + 1, cols=len(header))
            for col, cell in enumerate(header):
                table.cell(0, col).text = cell
            for r, row in enumerate(rows, start=1):
                for col, cell in enumerate(row):
                    table.cell(r, col).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return _normalize_zip(buffer.getvalue())


def _pdf_safe(text: str) -> str:
    from xml.sax.saxutils import escape

    # Keep "Version 2.0" / "rule 4.2" glued so a line never wraps to start with
    # a decimal that the reader would mistake for a subsection heading.
    return DECIMAL_AFTER_WORD.sub("\\1\u00a0\\2", escape(text))


def render_pdf(family, version) -> bytes:
    from reportlab import rl_config

    rl_config.invariant = 1
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table

    styles = getSampleStyleSheet()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        title=display_title(family, version),
        author="Doofenshmirtz Evil Inc Compliance (Norm)",
        creator="scripts/generate_corpus.py",
    )
    story = []
    for index, line in enumerate(header_lines(family, version)):
        style = styles["Title"] if index == 0 else styles["Heading2"]
        story.append(Paragraph(_pdf_safe(line), style))
    story.append(Spacer(1, 12))
    for kind, value in items(version["sections"]):
        if kind == "heading":
            story.append(Paragraph(_pdf_safe(value), styles["Heading3"]))
        elif kind in {"para", "bullet"}:
            text = value if kind == "para" else f"- {value}"
            story.append(Paragraph(_pdf_safe(text), styles["BodyText"]))
        else:
            table = Table(value)
            table.setStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey)])
            story.append(table)
            story.append(Spacer(1, 6))
    doc.build(story)
    return buffer.getvalue()


RENDERERS = {"md": render_md, "docx": render_docx, "pdf": render_pdf}


# ------------------------------------------------------------- manifest
def manifest_entry(family: dict, version: dict) -> dict:
    entry = {
        "department": family["department"],
        "doc_type": family.get("doc_type", "policy"),
        "classification": family.get("classification", "internal"),
        "status": version.get("status", family.get("status", "active")),
        "effective_from": version["date"],
    }
    return entry


def build(docs: Path, check: bool = False) -> dict:
    from rag.reader import read

    docs.mkdir(parents=True, exist_ok=True)
    manifest = {name: dict(entry) for name, entry in ORIGINALS.items()}
    written, stale, problems = 0, [], []
    names = set()
    for family in FAMILIES:
        for version in resolve(family):
            name = filename(family, version)
            if name in names or name in ORIGINALS:
                raise ValueError(f"duplicate document {name}")
            names.add(name)
            data = RENDERERS[version["format"]](family, version)
            path = docs / name
            if check:
                if not path.is_file() or path.read_bytes() != data:
                    stale.append(name)
            else:
                path.write_bytes(data)
                written += 1
                got = [b["heading"] for b in read(path)["blocks"]]
                want = expected_headings(version["sections"])
                if got != want:
                    problems.append((name, sorted(set(want) ^ set(got))[:4]))
            manifest[name] = manifest_entry(family, version)
    payload = (
        json.dumps(
            {
                "_comment": "Generated by scripts/generate_corpus.py. Per-document "
                "department, type, classification, lifecycle status and effective "
                "date. Edit the generator, not this file.",
                "documents": dict(sorted(manifest.items())),
            },
            indent=2,
            ensure_ascii=False,
        )
        + "\n"
    )
    manifest_path = docs / "manifest.json"
    if check:
        if not manifest_path.is_file() or manifest_path.read_text() != payload:
            stale.append("manifest.json")
    else:
        manifest_path.write_text(payload, encoding="utf-8")
    if problems:
        raise SystemExit(f"heading round-trip failed: {problems}")
    return {"written": written, "documents": len(manifest), "stale": stale}


def main(argv=None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docs", default=str(ROOT / "docs"))
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args(argv)
    result = build(Path(args.docs), check=args.check)
    if args.check:
        if result["stale"]:
            print(f"stale: {result['stale']}")
            return 1
        print(f"up to date: {result['documents']} documents")
        return 0
    print(f"wrote {result['written']} files; manifest has {result['documents']} docs")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
